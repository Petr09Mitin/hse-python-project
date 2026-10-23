import os
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from email.parser import BytesParser
from email.policy import default
from io import BytesIO
from urllib.parse import parse_qs, urlsplit
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from telebot import apihelper

from receipt_bot.application.messaging.incoming_message import IncomingMessage
from receipt_bot.bootstrap.application_settings import ApplicationSettings
from receipt_bot.bootstrap.build_application import build_application
from receipt_bot.domain.categories.category import DEFAULT_CATEGORIES
from receipt_bot.domain.money import Money
from tests.helpers.http.http_server import scripted_server
from tests.integration.test_qwen_adapters import completion
from tests.integration.test_telegram_adapters import telegram_message


def image_receipt(lines, image_format="PNG"):
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", 48)
    with Image.new("RGB", (1200, len(lines) * 100 + 100), "white") as picture, BytesIO() as output:
        draw = ImageDraw.Draw(picture)
        for index, line in enumerate(lines):
            draw.text((55, 40 + index * 100), line, font=font, fill="black")
        picture.save(output, format=image_format)
        return output.getvalue()


def receipt_file(kind):
    import pymupdf

    header = ("КАССОВЫЙ ЧЕК", "ООО МАГАЗИН", "22.10.2026")
    items = ("Молоко     89.90", "Хлеб       45.10", "ИТОГ      135.00")
    if kind == "photo":
        return "image/jpeg", image_receipt(header + items, "JPEG")
    with pymupdf.open() as document:
        page = document.new_page(width=1200, height=700)
        page.insert_font(fontname="receipt", fontfile="/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")
        lines = header + items if kind == "text_pdf" else header
        for index, line in enumerate(lines):
            page.insert_text((55, 85 + index * 100), line, fontsize=48, fontname="receipt")
        if kind == "mixed_pdf":
            page = document.new_page(width=600, height=200)
            page.insert_image(page.rect, stream=image_receipt(items))
        return "application/pdf", document.tobytes()


def settings(qwen_url, dsn, model_dir):
    return ApplicationSettings(
        {
            "TELEGRAM_API_KEY": "123456:test-token",
            "POSTGRES_DSN": dsn,
            "QWEN_URL": qwen_url,
            "MODEL_DIR": model_dir,
            "QWEN_TIMEOUT": "120",
            "OCR_TIMEOUT": "120",
        }
    )


def download_responses(content):
    metadata = {
        "ok": True,
        "result": {
            "file_id": "receipt-file",
            "file_unique_id": "receipt-unique",
            "file_size": len(content),
            "file_path": "receipt",
        },
    }
    return [(200, metadata), (200, content)]


def point_telegram(monkeypatch, url):
    monkeypatch.setattr(apihelper, "API_URL", url + "/bot{0}/{1}")
    monkeypatch.setattr(apihelper, "FILE_URL", url + "/file/bot{0}/{1}")


def received_xlsx(request):
    _, body, headers = request
    mime = ("Content-Type: " + headers["Content-Type"] + "\r\nMIME-Version: 1.0\r\n\r\n").encode() + body
    message = BytesParser(policy=default).parsebytes(mime)
    for part in message.walk():
        if part.get_filename() == "expenses.xlsx":
            return part.get_payload(decode=True)
    raise AssertionError("Telegram sendDocument did not contain expenses.xlsx")


@pytest.mark.parametrize("kind", ["photo", "text_pdf", "mixed_pdf"])
def test_assembled_real_recognition_save_restart_and_delivered_xlsx(kind, monkeypatch):
    """Проверяет полный путь чека и XLSX с настоящими моделями и БД; ответы Telegram заданы заранее."""
    from receipt_bot.adapters.persistence.schema import prepare_schema

    dsn = os.getenv("TEST_POSTGRES_DSN")
    qwen_url = os.getenv("TEST_QWEN_URL")
    model_dir = os.getenv("TEST_OCR_MODEL_DIR")
    if not all((dsn, qwen_url, model_dir)):
        pytest.skip("Full stack requires PostgreSQL, Qwen and Linux OCR models: make test")
    prepare_schema(dsn)
    owner = int(uuid4().int % 10**12) + 100
    media_type, content = receipt_file(kind)
    update_message = {
        **telegram_message(),
        "chat": {"id": owner, "type": "private"},
        "from": {"id": owner, "is_bot": False, "first_name": "QA"},
        "caption": "/receipt",
    }
    update_message.pop("text")
    if kind == "photo":
        update_message["photo"] = [
            {"file_id": "receipt-file", "file_unique_id": "receipt-unique", "width": 1200, "height": 700}
        ]
    else:
        update_message["document"] = {
            "file_id": "receipt-file",
            "file_unique_id": "receipt-unique",
            "file_name": "receipt.pdf",
            "mime_type": media_type,
            "file_size": len(content),
        }
    responses = [(200, {"ok": True, "result": [{"update_id": owner, "message": update_message}]})]
    responses += download_responses(content)
    responses += [(200, {"ok": True, "result": telegram_message()})] * 3
    app_settings = settings(qwen_url, dsn, model_dir)
    stages = []
    ocr_pages = []
    with scripted_server(responses) as (telegram, url):
        point_telegram(monkeypatch, url)
        runner = build_application(app_settings)
        client = runner.dispatcher.receipts.structurer.client
        original_request = client.request
        original_ocr = runner.dispatcher.receipts.extractor.ocr_runner.recognize

        def observe_request(name, schema, prompt):
            stages.append(name)
            return original_request(name, schema, prompt)

        def observe_ocr(image, page_number=1, page_width=None, page_height=None):
            ocr_pages.append(page_number)
            return original_ocr(image, page_number, page_width, page_height)

        monkeypatch.setattr(client, "request", observe_request)
        monkeypatch.setattr(runner.dispatcher.receipts.extractor.ocr_runner, "recognize", observe_ocr)
        try:
            incoming = runner.receiver.receive()
            assert incoming is not None and incoming.owner_id == owner
            outcomes = runner.dispatcher.dispatch(runner.dispatcher.prepare(incoming))
            assert all(outcome.status == "delivered" for outcome in outcomes)
            runner.receiver.acknowledge(incoming)
            stored = runner.dispatcher.receipts.repository.find_by_source(owner, incoming.attachments[0].source_key)
            assert stored is not None
            assert stored.receipt.is_reportable
            assert stored.receipt.purchase_date == date(2026, 10, 22)
            assert stored.receipt.total == Money(1350000)
            assert [(item.item_name, item.item_price) for item in stored.receipt.items] == [
                ("Молоко", Money(899000)),
                ("Хлеб", Money(451000)),
            ]
            sent = [request for request in telegram.requests if urlsplit(request[0]).path.endswith("/sendMessage")]
            summary = parse_qs(urlsplit(sent[0][0]).query)["text"][0]
            assert "135.0000" in summary
            assert "сохранён" in summary
        finally:
            runner.close()
    assert stages == ["receipt_decision", "receipt_data", "item_category", "item_category"]
    assert ocr_pages == ({"photo": [1], "text_pdf": [], "mixed_pdf": [2]}[kind])
    with scripted_server([(200, {"ok": True, "result": telegram_message()})] * 2) as (telegram, url):
        point_telegram(monkeypatch, url)
        restarted = build_application(app_settings)
        try:
            message = IncomingMessage(owner, owner, str(uuid4()), "/report 22.10.2026 22.10.2026")
            outcomes = restarted.dispatcher.dispatch(restarted.dispatcher.prepare(message))
            assert outcomes[0].status == "delivered"
            assert urlsplit(telegram.requests[-1][0]).path.endswith("/sendDocument")
            with BytesIO(received_xlsx(telegram.requests[-1])) as source:
                workbook = load_workbook(source, data_only=True)
                try:
                    sheet = workbook["Расходы"]
                    assert sheet.max_row == 4
                    assert sheet["A2"].value.date() == date(2026, 10, 22)
                    assert (sheet["B2"].value, sheet["C2"].value) == ("Молоко", 89.9)
                    assert (sheet["B3"].value, sheet["C3"].value) == ("Хлеб", 45.1)
                    assert sheet["C4"].value == 135
                    assert all(sheet.cell(row, 3).number_format == "0.0000" for row in (2, 3, 4))
                    assert workbook["Категории"].cell(workbook["Категории"].max_row, 2).value == 135
                finally:
                    workbook.close()
        finally:
            restarted.close()


def category_update(identifier, owner, task_id, category):
    return {
        "update_id": identifier,
        "callback_query": {
            "id": f"callback-{identifier}",
            "from": {"id": owner, "is_bot": False, "first_name": "QA"},
            "chat_instance": "qa",
            "message": {**telegram_message(), "chat": {"id": owner, "type": "private"}},
            "data": f"category:{task_id}:{category.code}",
        },
    }


def test_manual_review_race_restart_replay_and_four_decimal_delivered_report(monkeypatch):
    """Проверяет ручные категории, повторы и отчёт с настоящей БД и заданными внешними ответами."""
    from receipt_bot.adapters.persistence.schema import prepare_schema

    dsn, model_dir = os.getenv("TEST_POSTGRES_DSN"), os.getenv("TEST_OCR_MODEL_DIR")
    if not dsn or not model_dir:
        pytest.skip("Real PostgreSQL and model directories required: make test")
    prepare_schema(dsn)
    owner = int(uuid4().int % 10**12) + 100
    media_type, content = receipt_file("text_pdf")
    answers = [
        (200, completion({"is_receipt": True, "reason": "Receipt"})),
        (
            200,
            completion(
                {
                    "purchase_date": "2026-10-22",
                    "total": "12.3457",
                    "items": [
                        {"item_name": "Товар", "item_price": "12.3456"},
                        {"item_name": "Хлеб", "item_price": "0.0001"},
                    ],
                }
            ),
        ),
        (200, completion({"category_code": None, "needs_review": True})),
        (200, completion({"category_code": "food", "needs_review": False})),
    ]
    from receipt_bot.application.messaging.attachment_reference import AttachmentReference

    reference = AttachmentReference(str(uuid4()), 0, "receipt-file", "receipt.pdf", media_type)
    incoming = IncomingMessage(owner, owner, str(uuid4()), "/receipt", (reference,))
    success = (200, {"ok": True, "result": telegram_message()})
    callback_ack = (200, {"ok": True, "result": True})
    with scripted_server(answers) as (model, llm_url):
        app_settings = settings(llm_url + "/v1", dsn, model_dir)
        with scripted_server(download_responses(content) + [success] * 2) as (telegram, url):
            point_telegram(monkeypatch, url)
            runner = build_application(app_settings)
            try:
                assert all(
                    result.status == "delivered"
                    for result in runner.dispatcher.dispatch(runner.dispatcher.prepare(incoming))
                )
                stored = runner.dispatcher.receipts.repository.find_by_source(owner, reference.source_key)
                assert stored.receipt.is_reportable and stored.receipt.total == Money(123457)
                assert len(stored.review_tasks) == 1
                task = stored.review_tasks[0]
                markup = parse_qs(urlsplit(telegram.requests[-1][0]).query)["reply_markup"][0]
                assert f"category:{task.id}:food" in markup
            finally:
                runner.close()
        runner = build_application(app_settings)
        try:
            foreign = category_update(owner, owner + 1, task.id, DEFAULT_CATEGORIES[1])
            with scripted_server([(200, {"ok": True, "result": [foreign]}), callback_ack, success]) as (_, url):
                point_telegram(monkeypatch, url)
                message = runner.receiver.receive()
                assert runner.dispatcher.dispatch(runner.dispatcher.prepare(message))[0].status == "delivered"
                runner.receiver.acknowledge(message)
            repository = runner.dispatcher.receipts.repository
            assert repository.get_review_task(owner, task.id).status == "pending"
            choices = DEFAULT_CATEGORIES[:2]
            competing = [category_update(owner + i + 1, owner, task.id, choice) for i, choice in enumerate(choices)]
            with scripted_server(
                [(200, {"ok": True, "result": competing}), callback_ack, callback_ack, success, success]
            ) as (_, url):
                point_telegram(monkeypatch, url)
                messages = [runner.receiver.receive(), runner.receiver.receive()]
                prepared = [runner.dispatcher.prepare(message) for message in messages]
                with ThreadPoolExecutor(2) as workers:
                    results = list(workers.map(runner.dispatcher.dispatch, prepared))
                assert all(result[0].status == "delivered" for result in results)
                for message in messages:
                    runner.receiver.acknowledge(message)
            answered = repository.get_review_task(owner, task.id)
            assert answered.status == "answered" and answered.answer_category in choices
            winner = answered.answer_category
            replay = category_update(owner + 3, owner, task.id, DEFAULT_CATEGORIES[-1])
            with scripted_server([(200, {"ok": True, "result": [replay]}), callback_ack, success]) as (_, url):
                point_telegram(monkeypatch, url)
                message = runner.receiver.receive()
                runner.dispatcher.dispatch(runner.dispatcher.prepare(message))
                runner.receiver.acknowledge(message)
            assert repository.get_review_task(owner, task.id).answer_category == winner
        finally:
            runner.close()
        with scripted_server([success]) as (telegram, url):
            point_telegram(monkeypatch, url)
            restarted = build_application(settings("http://127.0.0.1:1/v1", dsn, model_dir))
            try:
                assert restarted.dispatcher.dispatch(restarted.dispatcher.prepare(incoming))[0].status == "delivered"
                assert len(telegram.requests) == 1
                assert urlsplit(telegram.requests[0][0]).path.endswith("/sendMessage")
                saved = restarted.dispatcher.receipts.repository.find_by_source(owner, reference.source_key)
                assert saved.receipt.items[0].category == winner
            finally:
                restarted.close()
        assert len(model.requests) == 4
        with scripted_server([success] * 2) as (telegram, url):
            point_telegram(monkeypatch, url)
            runner = build_application(app_settings)
            try:
                report = IncomingMessage(owner, owner, str(uuid4()), "/report 22.10.2026 22.10.2026")
                assert runner.dispatcher.dispatch(runner.dispatcher.prepare(report))[0].status == "delivered"
                with BytesIO(received_xlsx(telegram.requests[-1])) as source:
                    workbook = load_workbook(source, data_only=True)
                    try:
                        sheet = workbook["Расходы"]
                        assert sheet.max_row == 4
                        assert [Decimal(str(sheet.cell(row, 3).value)) for row in (2, 3, 4)] == [
                            Decimal("12.3456"),
                            Decimal("0.0001"),
                            Decimal("12.3457"),
                        ]
                        assert sheet["D2"].value == winner.name
                        assert all(sheet.cell(row, 3).number_format == "0.0000" for row in (2, 3, 4))
                        categories = workbook["Категории"]
                        assert Decimal(str(categories.cell(categories.max_row, 2).value)) == Decimal("12.3457")
                    finally:
                        workbook.close()
                foreign_report = IncomingMessage(owner + 1, owner + 1, str(uuid4()), "/report 22.10.2026 22.10.2026")
                prepared = runner.dispatcher.prepare(foreign_report)
                outgoing = runner.dispatcher.reports.process(foreign_report, prepared.resolution)
                assert outgoing.export_file is None
            finally:
                runner.close()
