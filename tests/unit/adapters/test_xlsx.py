from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from receipt_bot.domain.categories.category import DEFAULT_CATEGORIES
from receipt_bot.domain.money import Money
from receipt_bot.domain.reports.expense_report import ExpenseReport
from receipt_bot.domain.reports.expense_report_row import ExpenseReportRow
from receipt_bot.domain.reports.report_period import ReportPeriod


def test_xlsx_keeps_four_decimals_totals_and_untrusted_names_as_text():
    from receipt_bot.adapters.reports.openpyxl_report_exporter import OpenpyxlReportExporter

    report = ExpenseReport(
        ReportPeriod(date(2026, 10, 1), date(2026, 10, 22)),
        (
            ExpenseReportRow(date(2026, 10, 1), '=HYPERLINK("bad")', Money(123456), DEFAULT_CATEGORIES[0]),
            ExpenseReportRow(date(2026, 10, 22), "+cmd", Money(1), None),
        ),
    )
    exported = OpenpyxlReportExporter().export(report)
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook["Расходы"]
    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value == '=HYPERLINK("bad")'
    assert Decimal(str(sheet["C2"].value)) == Decimal("12.3456")
    assert sheet["C2"].number_format == "0.0000"
    assert sheet["C4"].value == 12.3457
    assert workbook["Категории"]["B2"].value == 12.3456
    assert workbook["Категории"]["A3"].value == "Требует проверки"
    workbook.close()


def test_xlsx_large_values_are_exact_strings_with_note():
    from receipt_bot.adapters.reports.openpyxl_report_exporter import OpenpyxlReportExporter

    report = ExpenseReport(
        ReportPeriod(date(2026, 1, 1), date(2026, 1, 1)),
        (ExpenseReportRow(date(2026, 1, 1), "Товар", Money(1234567890123456789), None),),
    )
    exported = OpenpyxlReportExporter().export(report)
    workbook = load_workbook(BytesIO(exported.content))
    assert workbook["Расходы"]["C2"].value == "123456789012345.6789"
    assert workbook["Расходы"]["C3"].value == "123456789012345.6789"
    assert "текст" in workbook["Примечание"]["A1"].value.lower()
    workbook.close()


@pytest.mark.parametrize("invalid_character", ["\x0b", "\x00", "\ud800", "\ufffe", "\uffff"])
def test_xlsx_normalizes_invalid_xml_characters_and_keeps_money_and_formula_safety(invalid_character):
    from receipt_bot.adapters.reports.openpyxl_report_exporter import OpenpyxlReportExporter
    from receipt_bot.domain.categories.category import Category

    category = Category("food", f"=Продукты{invalid_character}питание")
    report = ExpenseReport(
        ReportPeriod(date(2026, 10, 1), date(2026, 10, 1)),
        (
            ExpenseReportRow(date(2026, 10, 1), f"Молоко{invalid_character}1 л", Money(123456), category),
            ExpenseReportRow(date(2026, 10, 1), f"=SUM(1,2){invalid_character}товар", Money(1), None),
        ),
    )
    exported = OpenpyxlReportExporter().export(report)
    workbook = load_workbook(BytesIO(exported.content))
    try:
        sheet = workbook["Расходы"]
        assert sheet["B2"].value == "Молоко 1 л"
        assert sheet["B3"].value == "=SUM(1,2) товар"
        assert sheet["B3"].data_type == "s"
        assert sheet["D2"].value == "=Продукты питание"
        assert sheet["D2"].data_type == "s"
        assert Decimal(str(sheet["C2"].value)) == Decimal("12.3456")
        assert Decimal(str(sheet["C4"].value)) == Decimal("12.3457")
        assert workbook["Категории"]["A2"].value == "=Продукты питание"
        assert workbook["Категории"]["A2"].data_type == "s"
    finally:
        workbook.close()
