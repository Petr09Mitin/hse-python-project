import re
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook

from receipt_bot.application.messaging.attachment import Attachment
from receipt_bot.application.reports.report_exporter import ReportExporter
from receipt_bot.domain.money import Money
from receipt_bot.domain.reports.expense_report import ExpenseReport


def _xml_text(value: str) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", " ", value)


class OpenpyxlReportExporter(ReportExporter):
    """XLSX с заранее рассчитанными итогами и точными денежными строками."""

    def export(self, report: ExpenseReport) -> Attachment:
        workbook = Workbook()
        overflow = False
        try:
            expenses = workbook.active
            expenses.title = "Расходы"
            expenses.append(("Дата покупки", "Товар", "Сумма, ₽", "Категория"))
            for row in report.rows:
                expenses.append(
                    (
                        row.purchase_date,
                        _xml_text(row.item_name),
                        self._amount(row.price),
                        "Требует проверки" if row.category is None else _xml_text(row.category.name),
                    )
                )
                expenses.cell(expenses.max_row, 2).data_type = "s"
                expenses.cell(expenses.max_row, 4).data_type = "s"
                expenses.cell(expenses.max_row, 1).number_format = "DD.MM.YYYY"
            expenses.append(("Итого за период", "", self._amount(report.total)))
            categories = workbook.create_sheet("Категории")
            categories.append(("Категория", "Сумма, ₽"))
            for category, total in report.category_totals.items():
                categories.append(
                    ("Требует проверки" if category is None else _xml_text(category.name), self._amount(total))
                )
                categories.cell(categories.max_row, 1).data_type = "s"
            categories.append(("Итого за период", self._amount(report.total)))
            for sheet, column in ((expenses, 3), (categories, 2)):
                sheet.freeze_panes = "A2"
                for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
                    cell = cells[0]
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                        overflow = True
                    cell.number_format = "0.0000"
            if overflow:
                workbook.create_sheet("Примечание").append(
                    (
                        (
                            "Суммы свыше 15 значащих цифр записаны текстом для сохранения точности. "
                            "Итоги рассчитаны до экспорта."
                        ),
                    )
                )
            with BytesIO() as output:
                workbook.save(output)
                content = output.getvalue()
        finally:
            workbook.close()
        return Attachment("expenses.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", content)

    def _amount(self, money: Money):
        whole, fraction = divmod(money.units, 10000)
        exact = f"{whole}.{fraction:04d}"
        if len(str(money.units).rstrip("0")) > 15:
            return exact
        # Excel сохраняет не более 15 значащих цифр; большие суммы выше возвращаются текстом.
        return Decimal(exact)
